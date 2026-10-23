"""Skapa syntetiska exempelfiler som speglar typiska kodverk i kalkylark.

Innehållet är påhittat men strukturen följer verkliga fall:

* ``vardkontaktorsak.xlsx`` – ett enkelt kodverk med hierarki och extrakolumner
* ``kodverk_familj.xlsx`` – en familj av likadana flikar med nyckel–värde-huvud,
  grupprubriken "Används av" med kryss per användningskontext, härledda text-ID:n,
  fotnot, inledning med katalog och revisionshistorik, regler i löptext,
  en textkatalog och en hjälplista
* ``lab_urval.xlsx`` – urval ur NPU/SWE-koder: ett kort urval med grupp (och en
  felskriven kod), en längre grupperad lista med tomrader och en stor referenslista
* ``lista_utan_koder.pdf`` – tre listor med bara benämningar (ska avvisas)

Kör:  python examples/make_examples.py [målkatalog]
"""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook

CONTEXTS = ["Intyg A", "Intyg B", "Utlåtande C"]


def make_simple_workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kodverk"
    ws.append(["Kodverk för vårdkontaktens orsak, version 3"])
    ws.append([])
    ws.append(["Kod", "Benämning", "Beskrivning", "Överordnad kod", "Benämning (en)",
               "Synonymer", "Giltig fr.o.m.", "Inaktiv", "Anmärkning", "Sorteringsordning", "Ansvarig enhet", "Tom"])
    rows = [
        ("A", "Akut", "Oplanerad kontakt", None, "Acute", None, date(2020, 1, 1), "Nej", None, 1, "Enhet 1"),
        ("A1", "Akut sjukdom", "Akut insjuknande", "A", "Acute illness", "Akut insjuknande", date(2020, 1, 1), "Nej",
         None, 2, "Enhet 1"),
        ("A2", "Olycksfall", "Skada genom olycka", "A", "Accident", "Olycka;Trauma", date(2020, 1, 1), "Nej",
         "Se även B2", 3, "Enhet 2"),
        ("P", "Planerad", "Planerad kontakt", None, "Planned", None, date(2020, 1, 1), "Nej", None, 4, "Enhet 1"),
        ("P1", "Återbesök", "Planerat återbesök", "P", "Follow-up", "Uppföljning", date(2021, 3, 1), "Nej", None, 5,
         "Enhet 2"),
        ("P2", "Screening", "Kallelse till screening", "P", "Screening", None, date(2021, 3, 1), "Ja", "Utgår 2024", 6,
         "Enhet 2"),
    ]
    for r in rows:
        ws.append(list(r))
    wb.save(path)
    return path


def _kv_sheet(wb: Workbook, name: str, title: str, rows: list[tuple[str, str, list[int]]]) -> None:
    ws = wb.create_sheet(name)
    ws.append(["Kodverk:", name])
    ws.append(["Avser:", title])
    ws.append([])
    ws.append([None, None, None, None, None, "Används av"])
    ws.append(["Kod", "Klartext*", "ID rubriktext", "ID hjälptext", None] + CONTEXTS)
    for code, text, used in rows:
        ws.append([code, text, f"{name}.{code}.RBK", f"{name}.{code}.HLP", None]
                  + ["x" if i in used else None for i in range(len(CONTEXTS))])
    ws.append([])
    ws.append(["*Den exakta texten anges i textleveransen för respektive intyg."])


def make_family_workbook(path: Path) -> Path:
    wb = Workbook()
    intro = wb.active
    intro.title = "Inledning"
    intro.append([None, "Kodverk KV_EX"])
    intro.append([])
    intro.append([None, "Dokumentägare"])
    intro.append([None, "Exempelmyndigheten, enhet för kodverk"])
    intro.append([])
    intro.append([None, "Revisionshistorik"])
    intro.append([None, "Version", "Datum", "Beskrivning", None, None, "Namn"])
    intro.append([None, "1.0", date(2020, 1, 15), "Första version"])
    intro.append([None, "1.1", date(2021, 6, 1), "Lagt till KV_EX_0003"])
    intro.append([])
    intro.append([None, "Kodverk som ägs av myndigheten"])
    intro.append([None, "Namn", None, "OID/UUID", "Inkluderas", "Ägare/förvaltare", None, "Innehåll"])
    intro.append([None, "KV_EX_0001", None, "1.2.752.999.1.1", "ja", "Exempelmyndigheten", None, "Typ av underlag"])
    intro.append([None, "KV_EX_0002", None, "ej fastställt", "ja", "Exempelmyndigheten", None, "Typ av sysselsättning"])
    intro.append([None, "KV_EX_0003", None, "ej fastställt", "ja", "Exempelmyndigheten", None, "Grad"])

    about = wb.create_sheet("Om kodverk")
    about.append([None, "Om kodverk"])
    about.append([None, "Ett kodverk består av koder och visningsnamn. Visningsnamnen hanteras som texter och kan "
                        "uppdateras med kortare cykel. Det går däremot inte att göra förändringar som innebär att en "
                        "kod tas bort. Sådana förändringar hanteras genom urval. Det går bra att lägga till koder."])

    _kv_sheet(wb, "KV_EX_0001", "Typ av underlag", [
        ("UNDERSOKNING", "Min undersökning", [0, 1, 2]),
        ("TELEFON", "Telefonkontakt", [0]),
        ("JOURNAL", "Journaluppgifter", [0, 1]),
        ("ANNAT", "Annat", [0, 1, 2]),
    ])
    _kv_sheet(wb, "KV_EX_0002", "Typ av sysselsättning", [
        ("ARBETE", "Nuvarande arbete", [0, 2]),
        ("STUDIER", "Studier", [0]),
        ("ARBETSSOKANDE", "Arbetssökande", [0, 2]),
    ])
    _kv_sheet(wb, "KV_EX_0003", "Grad", [
        ("HELT", "100 procent", [1]),
        ("HALFTEN", "50 procent", [1]),
    ])
    helper = wb.create_sheet("datablad")
    helper.append([])
    helper.append([])
    helper.append([None, "Ja"])
    helper.append([None, "Nej"])

    texts = wb.create_sheet("Standardtexter")
    texts.append(["Standardtexter kodverk"])
    texts.append([])
    texts.append([None, "Texter"])
    texts.merge_cells("B3:C3")
    texts.append(["Kodverk\nUrval", "Rubrik-text", "Hjälp-text", "Rubriktexter", "Hjälptexter"])
    texts.append(["KV_EX_0001.UNDERSOKNING", "Ja", "Nej", "Min undersökning av patienten", None])
    texts.append(["KV_EX_0001.ANNAT", "Ja", "Ja", "Annat", "Ange vilket underlag som avses."])
    texts.append(["KV_EX_0002.ARBETE", "Ja", "Nej", "Nuvarande arbete", None])
    texts.append(["KV_EX_0002.SAKNAS", "Ja", "Nej", "Finns inte i kodverket", None])
    wb.save(path)
    return path


GROUPS = ["Hematologi", "Kemiska rutinanalyser", "Koagulation", "Urinanalyser"]


def make_lab_workbook(path: Path) -> Path:
    reference = [(f"NPU{10000 + i:05d}", f"P--Analys {i}") for i in range(600)]
    reference += [(f"SWE{5000 + i:05d}", f"S--Svensk analys {i}") for i in range(20)]
    wb = Workbook()
    short = wb.active
    short.title = "Urval"
    short.append(["Analys", "KOD", "Grupp"])
    short.append(["Pt-Kroppsvikt", "NPU1059", "Övrigt"])  # en siffra för lite; rätt kod finns nedan
    for i in range(1, 30):
        code, _ = reference[i * 7]
        short.append([f"P-Analys {i * 7}", code, GROUPS[i % len(GROUPS)]])
    short.append(["P-CRP", "SWE05003", "Övrigt"])

    grouped = wb.create_sheet("Grupperad")
    grouped.append(["Kortnamn", "Kod", "Kategori"])
    grouped.append(["Pt-Kroppsvikt", "NPU10599", "Övrigt"])
    for g, group in enumerate(GROUPS):
        for i in range(g * 40, g * 40 + 35):
            grouped.append([f"P-Analys {i}", reference[i][0], group])
        grouped.append([])
    grouped.append(["P-CRP", "SWE05003", "Övrigt"])

    ref = wb.create_sheet("Nomenklatur")
    ref.append(["Nomenklatur Kemi"])
    ref.append(["NPU/SWE-kod", "Trivialnamn"])
    for code, name in reference:
        ref.append([code, name])
    wb.save(path)
    return path


def make_list_pdf(path: Path) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    pages = [
        ["Vård och behandling", "Kvalitetssäkring", "Uppföljning och tillsyn av verksamheten", "Statistik"],
        ["Läsa", "Skriva", "Uppdatera", "Verifiera"],
        ["Vårdkontakt", "Vårdplan", "Diagnos", "Läkemedel ordination", "Undersökningsresultat"],
    ]
    c = canvas.Canvas(str(path), pagesize=A4)
    for lines in pages:
        y = 800
        for line in lines:
            c.setFont("Helvetica", 10)
            c.drawString(56, y, line)
            y -= 14
        c.showPage()
    c.save()
    return path


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).parent)
    for fn, name in [(make_simple_workbook, "vardkontaktorsak.xlsx"), (make_family_workbook, "kodverk_familj.xlsx"),
                     (make_lab_workbook, "lab_urval.xlsx"), (make_list_pdf, "lista_utan_koder.pdf")]:
        print("Skrev", fn(out / name))
