"""Läsning av Excel-/CSV-/PDF-filer till en gemensam bladmodell.

Kalkylark som används som kodverk har sällan en ren tabell. Läsaren tar
därför hand om:

* titel- och nyckel–värde-rader ovanför tabellen ("Kodverk: KV_X", "Avser: …")
  -> ``Sheet.metadata``
* tvånivårubriker, där en grupprubrik ("Används av") spänner över flera
  kolumner -> ``Sheet.column_groups``
* fotnoter under tabellen och fotnotsmarkörer i rubriker ("Klartext*")
  -> ``Sheet.footnotes`` och ``Sheet.header_notes``
* tomma avdelarrader och tomma avdelarkolumner
* PDF-utskrifter av kalkylark (se ``pdf.py``)
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

_MARKERS = ("*", "†", "¹", "²")


@dataclass
class Sheet:
    name: str
    header_row: int                     # 1-baserat radnummer för rubrikraden
    headers: list[str]
    rows: list[dict[str, Any]] = field(default_factory=list)
    total_rows: int = 0
    preamble: list[str] = field(default_factory=list)          # textrader ovanför rubriken
    metadata: dict[str, str] = field(default_factory=dict)     # "Kodverk:" -> "KV_X"
    column_groups: dict[str, list[str]] = field(default_factory=dict)
    footnotes: list[str] = field(default_factory=list)
    header_notes: dict[str, str] = field(default_factory=dict)  # rubrik -> fotnot
    raw: list[tuple] = field(default_factory=list, repr=False)  # alla rader, rensade värden
    source_format: str = "xlsx"

    def sample(self, n: int = 10) -> list[dict[str, Any]]:
        return self.rows[:n]

    def column(self, header: str) -> list[Any]:
        return [r.get(header) for r in self.rows]


def read_workbook(path: str | Path, sheet: str | None = None, max_rows: int | None = None) -> list[Sheet]:
    """Läs alla (eller ett namngivet) blad. ``max_rows`` begränsar antalet datarader."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".csv", ".tsv", ".txt"):
        sheets = [_read_csv(path, max_rows)]
    elif suffix == ".pdf":
        from .pdf import read_pdf
        sheets = read_pdf(path, max_rows)
    else:
        sheets = _read_xlsx(path, max_rows)
    if sheet:
        sheets = [s for s in sheets if s.name == sheet]
    return sheets


def _read_xlsx(path: Path, max_rows: int | None) -> list[Sheet]:
    wb = load_workbook(path, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        merged = {}
        for rng in ws.merged_cells.ranges:
            merged[(rng.min_row - 1, rng.min_col - 1)] = rng.max_col - 1
        raw = [tuple(_clean(v) for v in row) for row in ws.iter_rows(values_only=True)]
        parsed = parse_rows(ws.title, raw, max_rows, merged=merged)
        if parsed is not None:
            sheets.append(parsed)
    wb.close()
    return sheets


def _read_csv(path: Path, max_rows: int | None) -> Sheet:
    text = path.read_text(encoding="utf-8-sig")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        raw = [tuple(_clean(v) for v in r) for r in csv.reader(text.splitlines(), dialect)]
    except csv.Error:  # t.ex. en enda kolumn utan avgränsare
        raw = [tuple(_clean(v) for v in r) for r in csv.reader(text.splitlines(), delimiter=";")]
    sheet = parse_rows(path.stem, raw, max_rows) or Sheet(path.stem, 1, [])
    sheet.source_format = "csv"
    return sheet


def parse_rows(name: str, raw: list[tuple], max_rows: int | None = None,
               merged: dict[tuple[int, int], int] | None = None) -> Sheet | None:
    """Tolka rader (redan rensade värden) till ett ``Sheet``."""
    if not any(any(v is not None for v in r) for r in raw):
        return None
    header_idx = detect_header_row(raw)
    if header_idx < 0:
        # Ingen rubrikrad: en ren lista. Kolumnerna får generiska namn.
        width = max(len(r) for r in raw)
        used = [i for i in range(width) if any(i < len(r) and r[i] is not None for r in raw)]
        raw = [tuple(c for i, c in enumerate(r) if i in used) for r in raw]
        raw = [tuple(f"Kolumn {n + 1}" for n in range(len(used)))] + raw
        header_idx = 0
        synthetic_header = True
    else:
        synthetic_header = False
    headers, header_notes = _headers(raw[header_idx])
    header_cols = [i for i, h in enumerate(headers) if h]

    # Grupprubrik direkt ovanför rubrikraden ("Används av" över flera kolumner).
    groups: dict[str, list[str]] = {}
    group_row = header_idx - 1 if header_idx > 0 else None
    if group_row is not None:
        groups = _column_groups(raw[group_row], headers, header_cols, group_row, merged or {})
        if not groups:
            group_row = None

    preamble, metadata = [], {}
    for i, r in enumerate(raw[:header_idx]):
        if i == group_row:
            continue
        cells = [str(v) for v in r if v is not None]
        if not cells:
            continue
        if len(cells) >= 2 and cells[0].endswith(":"):
            metadata[cells[0].rstrip(":").strip()] = " ".join(cells[1:])
        else:
            preamble.append(" ".join(cells))

    rows, footnotes = [], []
    for r in raw[header_idx + 1:]:
        filled = [(i, v) for i, v in enumerate(r) if v is not None]
        if not filled:
            continue
        if len(header_cols) >= 2 and _is_footnote(filled):
            footnotes.append(str(filled[0][1]))
            continue
        rows.append({h: (r[i] if i < len(r) else None) for i, h in enumerate(headers) if h})

    # Fotnoter som börjar med en markör kopplas till rubriker med samma markör.
    for header, marker in list(header_notes.items()):
        note = next((f for f in footnotes if f.startswith(marker)), None)
        if note:
            header_notes[header] = note.lstrip("".join(_MARKERS)).strip()
        else:
            del header_notes[header]

    total = len(rows)
    if max_rows is not None:
        rows = rows[:max_rows]
    return Sheet(name=name, header_row=0 if synthetic_header else header_idx + 1,
                 headers=[h for h in headers if h], rows=rows,
                 total_rows=total, preamble=preamble, metadata=metadata, column_groups=groups,
                 footnotes=footnotes, header_notes=header_notes,
                 raw=raw[1:] if synthetic_header else raw)


def detect_header_row(raw: list[tuple], scan: int = 25) -> int:
    """Gissa rubrikraden: en rad med flera korta, unika texter, följd av data.

    Returnerar -1 om ingen rad ser ut som en rubrik (t.ex. en ren lista)."""
    best_idx, best_score = -1, -1.0
    width = max((len(r) for r in raw[:scan]), default=0)
    for i, row in enumerate(raw[:scan]):
        filled = [c for c in row if c is not None]
        if len(filled) < 2:
            continue
        texts = [c for c in filled if isinstance(c, str) and len(c) <= 60 and not c.endswith(":")]
        below = [r for r in raw[i + 1:i + 6] if any(v is not None for v in r)]
        below_filled = max((sum(v is not None for v in r) for r in below), default=0)
        score = len(texts) / max(width, 1)
        score += 0.3 if below and below_filled >= len(filled) * 0.5 else -0.3
        score -= 0.01 * i
        if len(set(texts)) < len(texts):
            score -= 0.2
        if score > best_score:
            best_idx, best_score = i, score
    return best_idx


def _headers(cells: tuple) -> tuple[list[str], dict[str, str]]:
    seen: dict[str, int] = {}
    out, notes = [], {}
    for cell in cells:
        h = re.sub(r"\s+", " ", str(cell)).strip() if cell is not None else ""
        marker = next((m for m in _MARKERS if h.endswith(m) and len(h) > 1), None)
        if marker:
            h = h[: -len(marker)].strip()
        if h:
            if h in seen:
                seen[h] += 1
                h = f"{h} ({seen[h]})"
            else:
                seen[h] = 1
            if marker:
                notes[h] = marker
        out.append(h)
    return out, notes


def _column_groups(row: tuple, headers: list[str], header_cols: list[int], row_idx: int,
                   merged: dict[tuple[int, int], int]) -> dict[str, list[str]]:
    filled = [(i, str(v)) for i, v in enumerate(row) if v is not None]
    if not filled or len(filled) > max(1, len(header_cols) // 2):
        return {}
    if any(i not in header_cols for i, _ in filled):
        return {}
    # En rubrik över första kolumnen är oftast en avsnittstitel ("Revisionshistorik"),
    # utom när cellen uttryckligen är sammanfogad över flera kolumner.
    if filled[0][0] == header_cols[0] and (row_idx, filled[0][0]) not in merged:
        return {}
    groups: dict[str, list[str]] = {}
    starts = [i for i, _ in filled]
    for n, (col, label) in enumerate(filled):
        if (row_idx, col) in merged:
            end = merged[(row_idx, col)]
        else:
            nxt = starts[n + 1] if n + 1 < len(starts) else len(headers)
            # Gruppen fortsätter så länge rubrikkolumnerna är sammanhängande.
            end = col
            while end + 1 < nxt and end + 1 < len(headers) and headers[end + 1]:
                end += 1
        members = [headers[i] for i in range(col, end + 1) if i < len(headers) and headers[i]]
        if len(members) >= 2:
            groups[label] = members
    return groups


def _is_footnote(filled: list[tuple[int, Any]]) -> bool:
    if len(filled) != 1:
        return False
    text = str(filled[0][1])
    return text.startswith(_MARKERS) or (len(text) > 80 and filled[0][0] == 0)


def _clean(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value
